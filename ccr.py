from os import path, system

DEFAULT_THEME_NAME = 'dracula'
DEFAULT_THEME_TONE = 'dark'

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:
    print('Auto Installing openpyxl')
    system('pip install openpyxl')
try:
    import ttkbootstrap as ttk
    from ttkbootstrap.constants import *
except ModuleNotFoundError:
    print('Auto Installing ttkbootstrap')
    system('pip install ttkbootstrap')
    import ttkbootstrap as ttk
    from ttkbootstrap.constants import *
    from openpyxl import Workbook, load_workbook, styles

def focus_next_widget(event):
    event.widget.tk_focusNext().focus()
    return("break")

root = path.dirname(__file__)
wb = load_workbook(f'{root}/lib/Template.xlsx')
ws = wb.active

row = 1

headings = []

while ws[f'A{row}'].value:
    headings.append(ws[f'A{row}'].value)
    row += 1

print(f'{DEFAULT_THEME_NAME}-{DEFAULT_THEME_TONE}')
window = ttk.App(theme=f'{DEFAULT_THEME_NAME}-{DEFAULT_THEME_TONE}')
window.tk.call('tk', 'scaling', 0.75)
window.title('Change Request Template')
window.columnconfigure(1, weight=1)
window.minsize(900, 600)

themes = ['bootstrap','pydata','nord','solarized','catppuccin','gruvbox','dracula','tokyo-night','one','everforest','vapor','minty','pulse','united','sandstone','']


theme_button = ttk.Button(master=window, text=f'Cycle Theme (Current: {window.style.theme_use()})', command=lambda: cycle_theme())
theme_button.grid(row=0, column=0, columnspan=2, sticky='nesw')


counter = 1
def cycle_theme():
    global counter
    if counter >= len(themes):
        counter = 0
    window.style.theme_use(themes[counter])
    theme_button.configure(text=f'Cycle Theme (Current: {window.style.theme_use()})')
    counter += 1

combos = {}

for heading in headings:
    combos[heading] = {}
    combos[heading]['row'] = headings.index(heading)+1
    combos[heading]['label'] = ttk.Label(master=window, text=heading, anchor='e')
    combos[heading]['label'].grid(row=combos[heading]['row'], column=0, sticky='nesw')
    if heading == 'Request for Change Details(Change Order, Request, or Incident Number)':
        combos[heading]['textentry'] = None
        continue
    combos[heading]['textentry'] = ttk.Text(master=window, height=2)
    combos[heading]['textentry'].bind("<Tab>", focus_next_widget)
    combos[heading]['textentry'].grid(row=combos[heading]['row'], column=1, sticky='nesw')

output_combo = {}
output_combo['heading'] = ttk.Label(master=window, text='Output File Name', anchor='e')
output_combo['heading'].grid(row=len(headings)+1, column=0, sticky='nesw')
output_combo['textentry'] = ttk.Text(master=window, height=2)
output_combo['textentry'].bind("<Tab>", focus_next_widget)
output_combo['textentry'].grid(row=len(headings)+1, column=1, sticky='nesw')

submit_button = ttk.Button(master=window, text='Submit', command=lambda : submit())
submit_button.grid(row=len(headings)+2, column=0, columnspan=2, sticky='nesw')


def submit():
    for heading in combos:
        thingy = combos[heading]['textentry']
        ws[f'B{combos[heading]["row"]}'] = str(thingy.get('1.0', END)).strip() if thingy else ""

    outfile_name = output_combo['textentry'].get('1.0', END).strip()
    
    wb.save(f'{root}/output/{outfile_name}.xlsx')

window.mainloop()
